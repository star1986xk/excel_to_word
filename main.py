from docx import Document
from openpyxl import *

excel_filename = '1.xlsx'
word_filename = '1.docx'
wb = load_workbook(excel_filename)
ws = wb.active
for row in list(ws.rows)[1:]:
    A = 0
    B = 0
    C = 0
    D = 0
    doc = Document(word_filename)
    tb = doc.tables
    tb_rows = tb[0].rows
    tb_rows[0].cells[1].text = row[1].value
    tb_rows[0].cells[5].text = row[2].value
    tb_rows[0].cells[8].text = row[3].value
    tb_rows[1].cells[3].text = row[4].value
    tb_rows[1].cells[9].text = row[5].value
    tb_rows[2].cells[3].text = row[6].value.strftime('%Y-%m-%d')
    tb_rows[2].cells[6].text = row[7].value.strftime('%Y-%m-%d')
    tb_rows[2].cells[10].text = row[8].value

    for i, n in enumerate(range(6, 21)):
        tb_rows[n].cells[6].text = '优秀' if row[n + i + 3].value >= 85 else (
            '良好' if row[n + i + 3].value >= 75 else ('及格' if row[n + i + 3].value >= 60 else '不及格'))
        tb_rows[n].cells[8].text = str(row[n + i + 3].value if row[n + i + 3].value else '')
        tb_rows[n].cells[10].text = ('及格' if row[n + i + 4].value >= 60 else '不及格') if row[n + i + 4].value else ''
        tb_rows[n].cells[11].text = str(row[n + i + 4].value if row[n + i + 4].value else '')
        tb_rows[n].cells[12].text = tb_rows[n].cells[10].text if tb_rows[n].cells[10].text else tb_rows[n].cells[6].text
        if tb_rows[n].cells[12].text == '优秀':
            A +=1
        if tb_rows[n].cells[12].text == '良好':
            B += 1
        if tb_rows[n].cells[12].text == '及格':
            C +=1
        if tb_rows[n].cells[12].text == '不及格':
            D +=1
    tb_rows[3].cells[3].text = '课目优秀 {} 个，良好 {} 个，及格 {} 个，不及格 {} 个。'.format(A,B,C,D)
    tb_rows[3].cells[12].text = '不及格' if D else '及格'

    doc.save(row[1].value + ".docx")

